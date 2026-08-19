"""Render a MeetingExport to Markdown, JSON, PDF, or DOCX."""

from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass

from packages.export.models import ExportSummary, MeetingExport, format_timestamp

__all__ = ["FORMATS", "ExportFormat", "to_docx", "to_json", "to_markdown", "to_pdf"]


def _summary_sections(summary: ExportSummary) -> list[tuple[str, list[str]]]:
    return [
        ("Key points", summary.key_points),
        ("Action items", summary.action_items),
        ("Decisions", summary.decisions),
        ("Open questions", summary.open_questions),
    ]


def to_markdown(export: MeetingExport) -> str:
    lines = [f"# {export.title}", "", f"*{export.started_at} · {export.source}*", ""]

    if export.summary:
        lines += ["## Summary", "", export.summary.summary, ""]
        for heading, items in _summary_sections(export.summary):
            if not items:
                continue
            lines.append(f"### {heading}")
            checkbox = heading == "Action items"
            for item in items:
                lines.append(f"- [ ] {item}" if checkbox else f"- {item}")
            lines.append("")

    lines.append("## Transcript")
    lines.append("")
    for segment in export.segments:
        stamp = format_timestamp(segment.start_ms)
        who = f"**{segment.speaker}:** " if segment.speaker else ""
        lines.append(f"`[{stamp}]` {who}{segment.text}")
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def to_json(export: MeetingExport) -> str:
    return export.model_dump_json(indent=2)


def to_pdf(export: MeetingExport) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [
        Paragraph(_escape(export.title), styles["Title"]),
        Paragraph(f"{_escape(export.started_at)} · {_escape(export.source)}", styles["Italic"]),
        Spacer(1, 12),
    ]

    if export.summary:
        story.append(Paragraph("Summary", styles["Heading2"]))
        story.append(Paragraph(_escape(export.summary.summary), styles["BodyText"]))
        for heading, items in _summary_sections(export.summary):
            if not items:
                continue
            story.append(Paragraph(heading, styles["Heading3"]))
            for item in items:
                story.append(Paragraph(f"• {_escape(item)}", styles["BodyText"]))
        story.append(Spacer(1, 12))

    story.append(Paragraph("Transcript", styles["Heading2"]))
    for segment in export.segments:
        stamp = format_timestamp(segment.start_ms)
        who = f"<b>{_escape(segment.speaker)}:</b> " if segment.speaker else ""
        story.append(Paragraph(f"[{stamp}] {who}{_escape(segment.text)}", styles["BodyText"]))

    buffer = io.BytesIO()
    SimpleDocTemplate(buffer, pagesize=letter, title=export.title).build(story)
    return buffer.getvalue()


def to_docx(export: MeetingExport) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(export.title, level=0)
    document.add_paragraph(f"{export.started_at} · {export.source}").italic = True

    if export.summary:
        document.add_heading("Summary", level=1)
        document.add_paragraph(export.summary.summary)
        for heading, items in _summary_sections(export.summary):
            if not items:
                continue
            document.add_heading(heading, level=2)
            style = "List Bullet"
            for item in items:
                document.add_paragraph(item, style=style)

    document.add_heading("Transcript", level=1)
    for segment in export.segments:
        stamp = format_timestamp(segment.start_ms)
        paragraph = document.add_paragraph()
        paragraph.add_run(f"[{stamp}] ").italic = True
        if segment.speaker:
            paragraph.add_run(f"{segment.speaker}: ").bold = True
        paragraph.add_run(segment.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _escape(text: str) -> str:
    """Escape the reportlab mini-markup special characters."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


@dataclass
class ExportFormat:
    ext: str
    mime: str
    render: Callable[[MeetingExport], bytes]


FORMATS: dict[str, ExportFormat] = {
    "markdown": ExportFormat("md", "text/markdown", lambda e: to_markdown(e).encode()),
    "json": ExportFormat("json", "application/json", lambda e: to_json(e).encode()),
    "pdf": ExportFormat("pdf", "application/pdf", to_pdf),
    "docx": ExportFormat(
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        to_docx,
    ),
}
